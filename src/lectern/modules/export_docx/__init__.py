"""Export · DOCX: a Document as a Word file, with python-docx. Tables are tables; an open
question is its text with a comment listing the candidates; page notes, when asked for,
are comments on the page's first block."""

from __future__ import annotations

import io

import docx
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

from lectern.agent.tools import AgentTool
from lectern.documents.document import Block, Document
from lectern.documents.export import (
    ExportOptions,
    doubt_comment,
    figure_placeholder,
    inline_doubt,
    note_text,
    reading_order,
    segments,
)
from lectern.modules import ToolEnv

AUTHOR = "Lectern"


class DocxExporter:
    format = "docx"
    title = "Export · DOCX"
    extension = ".docx"
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    capabilities = frozenset({"tables", "comments"})

    def export(self, doc: Document, opts: ExportOptions) -> bytes:
        out = docx.Document()
        out.core_properties.title = doc.title
        for block, note, figure in reading_order(doc):
            first = _block(out, block, figure, opts)
            if note is not None and opts.notes == "comment" and first is not None:
                out.add_comment(first, text=note_text(block.page, note), author=AUTHOR)
        buffer = io.BytesIO()
        out.save(buffer)
        return buffer.getvalue()


def _runs(out: DocxDocument, paragraph: Paragraph, text: str, block: Block, opts: ExportOptions):  # noqa: ANN202
    """Write `text` into the paragraph; returns its first run (where a note can attach)."""
    first = None
    for run_text, doubt in segments(text, block.doubts) or [("", None)]:
        if doubt is None:
            run = paragraph.add_run(run_text)
        elif opts.open_doubts == "comment":
            run = paragraph.add_run(doubt.raw)
            out.add_comment(run, text=doubt_comment(doubt), author=AUTHOR)
        else:
            run = paragraph.add_run(inline_doubt(doubt))
        first = first or run
    return first


def _block(out: DocxDocument, block: Block, figure: int, opts: ExportOptions):  # noqa: ANN202
    if block.type == "heading":
        return _runs(out, out.add_heading(level=min(block.level or 2, 9)), block.text, block, opts)
    if block.type == "table":
        rows = block.rows or []
        if not rows:
            return None
        width = max(len(r) for r in rows)
        table = out.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        first = None
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                paragraph = table.cell(r, c).paragraphs[0]
                run = _runs(out, paragraph, cell, block, opts)
                if r == 0 and run is not None:
                    run.bold = True
                first = first or run
        return first
    if block.type == "list":
        first = None
        for item in (i for i in block.text.splitlines() if i.strip()):
            run = _runs(out, out.add_paragraph(style="List Bullet"), item, block, opts)
            first = first or run
        return first
    if block.type == "figure":
        return out.add_paragraph().add_run(figure_placeholder(block, figure))
    return _runs(out, out.add_paragraph(), block.text, block, opts)


class DocxModule:
    id = "export-docx"

    def tools(self, env: ToolEnv) -> list[AgentTool]:
        return []

    def exporters(self) -> list[DocxExporter]:
        return [DocxExporter()]


module = DocxModule()

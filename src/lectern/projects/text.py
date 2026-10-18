"""Text from a project file, for the chat to read and for search.

Plain text, Markdown and code are read as they are. PDFs give their text layer (a scan
without one has no text until it goes through OCR), DOCX its paragraphs and tables. The text
comes back in chunks: one per Markdown section, PDF page or ~1,500 characters, each with a
block reference `#c<n>` that search results and `document.read` use.
"""

from __future__ import annotations

import mimetypes
import re
import threading
from dataclasses import dataclass, field
from pathlib import Path

TEXT_LIMIT = 2 * 1024 * 1024  # characters of text kept per file
CHUNK_CHARS = 1500
# A PDF page with fewer characters than this is treated as having no text layer.
PDF_PAGE_MIN_CHARS = 20

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml", ".toml",
    ".xml", ".html", ".htm", ".ini", ".cfg", ".log", ".py", ".js", ".ts", ".tsx", ".jsx",
    ".css", ".sql", ".sh", ".ps1", ".bat", ".java", ".kt", ".go", ".rs", ".c", ".h", ".cpp",
    ".hpp", ".cs", ".php", ".rb", ".swift", ".tex",
}  # fmt: skip
# One PDFium call at a time, process-wide. The OCR pipeline (M2) must take it too.
PDFIUM_LOCK = threading.Lock()
MARKDOWN_EXTENSIONS = {".md", ".markdown"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_OVERRIDES = {".md": "text/markdown", ".markdown": "text/markdown", ".docx": DOCX_MIME}


# A Document's block anchor in document.md: {#b3.2 p=3}, {#b3.3 p=3→4}, {#b3.4 p=3 figure}
ANCHOR = re.compile(r"\{#((?:b|p)\d+(?:\.\d+)?) p=[\d→]+(?: [a-z]+)?\}")
NOTE = re.compile(r"^> Note for page (\d+)")
PAGE = re.compile(r"p=(\d+)")


@dataclass(frozen=True)
class Chunk:
    number: int  # 1-based; an ordinary file's block reference is #c<number>
    heading: str  # a Markdown heading, "Page 3", or ""
    body: str
    label: str | None = None  # a Document's block id ("b3.2"), cited as #b3.2 instead

    @property
    def ref(self) -> str:
        return f"#{self.label}" if self.label else f"#c{self.number}"


@dataclass
class Extracted:
    chunks: list[Chunk] = field(default_factory=list)
    note: str | None = None  # why there's no text, or that it was cut

    @property
    def has_text(self) -> bool:
        return bool(self.chunks)


def mime_for(name: str) -> str:
    suffix = Path(name).suffix.lower()
    if suffix in MIME_OVERRIDES:
        return MIME_OVERRIDES[suffix]
    guessed, _ = mimetypes.guess_type(name)
    if guessed:
        return guessed
    return "text/plain" if suffix in TEXT_EXTENSIONS else "application/octet-stream"


def extract(path: Path) -> Extracted:
    """The file's text in chunks, or a note saying why there is none."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf(path)
    if suffix == ".docx":
        return _limited(_docx(path), markdown=True)
    if suffix in TEXT_EXTENSIONS or mime_for(path.name).startswith("text/"):
        text = _read_text(path)
        if text is None:
            return Extracted(note="not a text file")
        if suffix in MARKDOWN_EXTENSIONS and ANCHOR.search(text):
            return Extracted(chunks=_document_chunks(text[:TEXT_LIMIT]))
        return _limited(text, markdown=suffix in MARKDOWN_EXTENSIONS)
    return Extracted(note="no text to read in this kind of file")


def _read_text(path: Path) -> str | None:
    with path.open("rb") as f:
        raw = f.read(TEXT_LIMIT * 4 + 1)
    if b"\x00" in raw[:8192]:
        return None  # binary
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _limited(text: str, *, markdown: bool) -> Extracted:
    note = None
    if len(text) > TEXT_LIMIT:
        text = text[:TEXT_LIMIT]
        note = f"only the first {TEXT_LIMIT // (1024 * 1024)} MB of text is used"
    sections = _markdown_sections(text) if markdown else [("", text)]
    return Extracted(chunks=_number(_split(sections)), note=note)


def _document_chunks(text: str) -> list[Chunk]:
    """A Document's document.md, one chunk per block, labelled with the block's id so the
    chat reads and cites it as #b3.2. What comes before the first block (the front matter) is
    a chunk of its own, and each page note is one too (#note3). A block's section is the last
    heading before it, or its page while there is none (a small model types few headings)."""
    pieces: list[tuple[str, str, str | None]] = []  # heading, body, label
    heading, label, lines = "About this document", None, []
    titled = False  # a heading block seen yet

    def flush() -> None:
        if label or "".join(lines).strip():
            pieces.append((heading, "\n".join(lines).strip(), label))

    for line in text.splitlines():
        anchor = ANCHOR.search(line)
        note = NOTE.match(line)
        if anchor and line.lstrip().startswith(("{#", "#")):
            flush()
            label = anchor.group(1)
            rest = ANCHOR.sub("", line).strip()
            if rest.startswith("#"):
                heading, titled = rest.lstrip("#").strip(), True
            elif not titled:
                heading = f"Page {PAGE.search(anchor.group(0)).group(1)}"
            lines = [rest] if rest else []
        elif note:
            flush()
            label, lines = f"note{note.group(1)}", [line]
            if not titled:
                heading = f"Page {note.group(1)}"
        else:
            lines.append(line)
    flush()
    return [Chunk(i, h, b, lb) for i, (h, b, lb) in enumerate(pieces, start=1)]


def _markdown_sections(text: str) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = []
    heading, lines = "", []
    for line in text.splitlines():
        match = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", line)
        if match:
            if "".join(lines).strip() or heading:
                sections.append((heading, "\n".join(lines)))
            heading, lines = match.group(1), []
        else:
            lines.append(line)
    if "".join(lines).strip() or heading:
        sections.append((heading, "\n".join(lines)))
    return sections


def _split(sections: list[tuple[str, str]]) -> list[tuple[str, str]]:
    """Sections longer than CHUNK_CHARS are cut on paragraph boundaries (or hard, for a
    paragraph that is itself too long)."""
    pieces: list[tuple[str, str]] = []
    for heading, body in sections:
        current = ""
        for paragraph in re.split(r"\n\s*\n", body.strip()):
            while len(paragraph) > CHUNK_CHARS:
                if current:
                    pieces.append((heading, current))
                    current = ""
                pieces.append((heading, paragraph[:CHUNK_CHARS]))
                paragraph = paragraph[CHUNK_CHARS:]
            if current and len(current) + len(paragraph) + 2 > CHUNK_CHARS:
                pieces.append((heading, current))
                current = ""
            current = f"{current}\n\n{paragraph}" if current else paragraph
        if current or heading:
            pieces.append((heading, current))
    return pieces


def _number(pieces: list[tuple[str, str]]) -> list[Chunk]:
    return [Chunk(i, heading, body) for i, (heading, body) in enumerate(pieces, start=1)]


def _pdf(path: Path) -> Extracted:
    # PDFium isn't thread-safe (pypdfium2's own docs), and imports run in worker threads.
    with PDFIUM_LOCK:
        return _pdf_text(path)


def pdf_pages(path: Path) -> list[str]:
    """Every page's text, "" for a page without any, read the same way as at import but all
    of it: TEXT_LIMIT bounds what the chat reads, and an export cut short would lose pages
    without saying so. Raises ValueError when the PDF can't be opened."""
    with PDFIUM_LOCK:
        read = _read_pages(path, limit=None)
    if read is None:
        raise ValueError(f"{path.name}: the PDF couldn't be opened")
    return read[0]


def scanned_pages(path: Path) -> int:
    """How many of a PDF's pages are scans among typed pages: no text layer, and an image on
    them. A blank page has neither, and nothing for OCR to read. 0 when it can't be opened."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c

    with PDFIUM_LOCK:
        try:
            pdf = pdfium.PdfDocument(path)
        except pdfium.PdfiumError:
            return 0
        try:
            scans = 0
            for index in range(len(pdf)):
                page = pdf[index]
                images = page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE])
                if len(_page_text(page)) < PDF_PAGE_MIN_CHARS and next(images, None) is not None:
                    scans += 1
                page.close()
            return scans
        finally:
            pdf.close()


def _read_pages(path: Path, limit: int | None) -> tuple[list[str], int] | None:
    """With PDFIUM_LOCK held. The pages' texts (stopping once past `limit` characters, if
    given) and the page count, or None when the PDF can't be opened."""
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(path)
    except pdfium.PdfiumError:
        return None
    try:
        pages: list[str] = []
        total = 0
        for index in range(len(pdf)):
            page = pdf[index]
            text = _page_text(page)
            page.close()
            pages.append(text)
            total += len(text)
            if limit is not None and total > limit:
                break
        return pages, len(pdf)
    finally:
        pdf.close()


def _page_text(page) -> str:  # noqa: ANN001 (a pypdfium2 PdfPage, imported lazily)
    textpage = page.get_textpage()
    text = textpage.get_text_bounded().replace("\r\n", "\n").strip()
    textpage.close()
    return text


def _pdf_text(path: Path) -> Extracted:
    read = _read_pages(path, limit=TEXT_LIMIT)
    if read is None:
        return Extracted(note="the PDF couldn't be opened")
    pages, count = read
    if sum(len(t) >= PDF_PAGE_MIN_CHARS for t in pages) == 0:
        return Extracted(note="no text layer (a scan: open it as an OCR session)")
    numbered = [(f"Page {n}", t) for n, t in enumerate(pages, start=1) if t]
    chunks = _number(_split(numbered))
    note = None
    if len(pages) < count:
        note = f"only the first {len(pages)} of {count} pages are used"
    return Extracted(chunks=chunks, note=note)


def docx_markdown(path: Path) -> str:
    """A DOCX as Markdown with real pipe tables, for building a Document from it (export)."""
    return _docx(path, pipe_tables=True)


def _docx(path: Path, *, pipe_tables: bool = False) -> str:
    """Paragraphs and tables in document order, headings as Markdown headings. Table rows are
    "a | b" for the search index; with pipe_tables, Markdown tables a parser recognises."""
    import docx
    from docx.table import Table

    document = docx.Document(str(path))
    lines: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            for n, row in enumerate(item.rows):
                cells = [cell.text.strip() for cell in row.cells]
                if not pipe_tables:
                    lines.append(" | ".join(cells))
                    continue
                safe = (c.replace("|", "\\|").replace("\n", " ") for c in cells)
                lines.append("| " + " | ".join(safe) + " |")
                if n == 0:
                    lines.append("| " + " | ".join("---" for _ in cells) + " |")
            lines.append("")
            continue
        text = item.text.strip()
        style = (item.style.name if item.style is not None else "") or ""
        level = re.match(r"Heading (\d)", style)
        if level and text:
            lines.append(f"{'#' * int(level.group(1))} {text}")
        elif style == "Title" and text:
            lines.append(f"# {text}")
        else:
            lines.append(text)
        lines.append("")
    return "\n".join(lines)

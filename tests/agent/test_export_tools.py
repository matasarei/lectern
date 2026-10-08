from __future__ import annotations

import asyncio
import inspect

import pytest

from lectern.agent.export_tools import export_tools
from lectern.agent.tools import ToolError
from lectern.core.auth import User
from lectern.modules import load_modules, module_exporters
from lectern.projects.files import Origin
from lectern.projects.tables import ArtifactRow

from ..documents.fixtures import lease
from .conftest import add_file

LETTER = "# Відповідь\n\n| Місяць | Сума |\n|---|---|\n| березень | 250,00 |\n"


def _tools():  # noqa: ANN202
    formats, run = export_tools(module_exporters(load_modules()))
    return formats, run


def _call(tool, ctx, **arguments):  # noqa: ANN001, ANN202
    result = tool.run(ctx, tool.validate(arguments))
    return asyncio.run(result) if inspect.isawaitable(result) else result


def _session_document(ctx) -> int:  # noqa: ANN001
    """An OCR session's output: Lease 1998/document.json and its document.md artifact."""
    folder = ctx.project.root / "Lease 1998"
    folder.mkdir()
    (folder / "document.json").write_text(lease().to_json(), encoding="utf-8")
    return add_file(ctx, "document.md", "{#b3.1 p=3} …", Origin("session", 9), folder="Lease 1998")


def _row(ctx, artifact_id: int) -> ArtifactRow:  # noqa: ANN001
    with ctx.sessions() as session:
        return session.get(ArtifactRow, artifact_id)


def test_formats_lists_each_installed_format_with_what_it_supports(ctx):
    formats, _ = _tools()
    text = _call(formats, ctx).text
    assert (
        "md · tables" in text
        and "docx · tables, comments" in text
        and "odt · tables, comments" in text
    )
    assert "open_doubts" in text and "notes" in text


def test_a_documents_folder_becomes_a_new_file_each_time_from_its_document(ctx):
    _, run = _tools()
    source = _session_document(ctx)

    first = _call(run, ctx, name="Lease 1998/", format="odt")
    again = _call(run, ctx, name="Lease 1998", format="odt")

    assert (
        first.refs[0]["name"] == "Lease 1998.odt" and again.refs[0]["name"] == "Lease 1998 (2).odt"
    )
    row = _row(ctx, first.refs[0]["id"])
    assert (row.from_kind, row.from_id, row.kind) == ("artifact", source, "generated")
    assert (ctx.project.root / "Lease 1998.odt").read_bytes()[:2] == b"PK"
    assert first.action == {"type": "open_tab", "tab": {"kind": "artifact", "id": row.id}}
    assert "1 open question, as a comment" in first.text


def test_a_markdown_file_goes_through_the_same_exporter(ctx):
    _, run = _tools()
    add_file(ctx, "reply-letter.md", LETTER)
    result = _call(run, ctx, name="reply-letter.md", format="docx")
    assert result.refs[0]["name"] == "reply-letter.docx"


def test_an_unknown_format_names_the_installed_ones_and_a_scan_says_what_to_do(ctx):
    _, run = _tools()
    add_file(ctx, "reply-letter.md", LETTER)
    with pytest.raises(ToolError, match="No pdf exporter is installed. Installed: docx, md, odt."):
        _call(run, ctx, name="reply-letter.md", format="pdf")
    add_file(ctx, "scan.png", "x")
    with pytest.raises(ToolError, match="isn't a Document, a PDF with text, a DOCX or a text"):
        _call(run, ctx, name="scan.png", format="odt")


def test_it_asks_first_only_when_the_format_would_lose_something(ctx):
    _, run = _tools()
    _session_document(ctx)
    asked: list[str] = []

    async def no(question: str) -> bool:
        asked.append(question)
        return False

    ctx.confirm = no
    _call(run, ctx, name="Lease 1998/", format="odt", open_doubts="comment")  # keeps them
    assert asked == []
    with pytest.raises(ToolError, match="not exported"):
        _call(run, ctx, name="Lease 1998/", format="md", open_doubts="comment")
    assert asked == ["Export Lease 1998 to md? It has no comments: the open questions become "
                     "inline marks."]  # fmt: skip


def test_someone_elses_file_isnt_found(ctx, store):
    _, run = _tools()
    add_file(ctx, "reply-letter.md", LETTER)
    ctx.user = User(id=2, name="other")
    with pytest.raises(ToolError, match="No file called"):
        _call(run, ctx, name="reply-letter.md", format="md")


def _add_bytes(ctx, name: str, data: bytes) -> None:  # noqa: ANN001
    from lectern.projects import files

    with ctx.sessions.begin() as session:
        files.write_bytes(
            ctx.project.root, session, owner_id=ctx.user.id, name=name, data=data,
            origin=Origin("import"),
        )  # fmt: skip


def test_a_pdf_with_text_exports_and_says_how_to_keep_its_layout(ctx, tmp_path):
    import docx

    from ..projects.conftest import write_pdf

    _, run = _tools()
    pdf = write_pdf(tmp_path / "s.pdf", ["Opening balance 100,00", "", "Closing balance 87,50"])
    _add_bytes(ctx, "statement.pdf", pdf.read_bytes())

    result = _call(run, ctx, name="statement.pdf", format="docx")

    assert result.refs[0]["name"] == "statement.docx"
    paragraphs = [
        p.text for p in docx.Document(str(ctx.project.root / "statement.docx")).paragraphs
    ]
    assert "Opening balance 100,00" in paragraphs and "Closing balance 87,50" in paragraphs
    assert "text layer" in result.text and "keep_layout" in result.text
    assert "Page 2 has little or no text (maybe a scan)" in result.text  # only OCR reads it


def test_a_docx_exports_to_odt_with_its_table(ctx, tmp_path):
    import zipfile

    from ..projects.conftest import write_docx

    _, run = _tools()
    _add_bytes(ctx, "lease.docx", write_docx(tmp_path / "l.docx").read_bytes())

    result = _call(run, ctx, name="lease.docx", format="odt")

    content = zipfile.ZipFile(ctx.project.root / "lease.odt").read("content.xml").decode()
    table = content[content.index("<table:table ") : content.index("</table:table>")]
    assert "250,00 грн" in table and "0,5% на день" in table
    assert "keep_layout" not in result.text  # only a PDF's text layer loses its layout


def test_a_scanned_page_with_only_a_page_number_still_counts_as_a_scan(ctx, tmp_path):
    from ..projects.conftest import write_pdf

    _, run = _tools()
    pages = ["Opening balance 100,00", "3", "", "Closing balance 87,50"]
    _add_bytes(ctx, "statement.pdf", write_pdf(tmp_path / "s.pdf", pages).read_bytes())

    result = _call(run, ctx, name="statement.pdf", format="md")

    assert "Pages 2, 3 have little or no text (maybe scans)" in result.text

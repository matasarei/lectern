from __future__ import annotations

from pathlib import Path


def write_pdf(path: Path, pages: list[str]) -> Path:
    """A minimal PDF with a text layer: one Helvetica line per page (Latin text only)."""
    objects: list[bytes] = []
    count = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(count))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode())
    font_id = 3 + 2 * count
    for i, text in enumerate(pages):
        content_id = 4 + 2 * i
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    size = len(objects) + 1
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref)
    path.write_bytes(bytes(out))
    return path


def write_scan_pdf(path: Path, pages: int = 2) -> Path:
    """A PDF with blank pages and no text layer, standing in for a scan."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument.new()
    for _ in range(pages):
        pdf.new_page(595, 842)
    pdf.save(path)
    pdf.close()
    return path


def add_pages(path: Path, *, blank: int = 0, scanned: int = 0) -> Path:
    """Pages added to the end of a PDF: blank ones (nothing on them), and scanned ones (one
    image each, no text layer)."""
    import pypdfium2 as pdfium
    from PIL import Image

    pdf = pdfium.PdfDocument(path)
    for _ in range(blank):
        pdf.new_page(612, 792)
    for _ in range(scanned):
        page = pdf.new_page(612, 792)
        image = pdfium.PdfImage.new(pdf)
        image.set_bitmap(pdfium.PdfBitmap.from_pil(Image.new("RGB", (40, 40), "gray")))
        image.set_matrix(pdfium.PdfMatrix().scale(400, 400))
        page.insert_obj(image)
        page.gen_content()
    out = path.with_suffix(".pages.pdf")  # pdfium can't save over the file it has open
    pdf.save(out)
    pdf.close()
    return out.replace(path)


def write_docx(path: Path) -> Path:
    import docx

    document = docx.Document()
    document.add_heading("Договір оренди", level=1)
    document.add_paragraph("Орендар сплачує пеню за кожен день прострочення.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Сума"
    table.cell(0, 1).text = "250,00 грн"
    table.cell(1, 0).text = "Пеня"
    table.cell(1, 1).text = "0,5% на день"
    document.add_heading("Підписи", level=2)
    document.add_paragraph("Ковальчук О. П.")
    document.save(str(path))
    return path
